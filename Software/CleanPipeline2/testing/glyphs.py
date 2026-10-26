"""Every stored glyph of every author, drawn, with its chance of being used.

Writes results/Glyphs.docx and results/glyphs.json

The probabilities are not stored in the profile, they come out of the selection
rule in synthesise._GlyphSource, recomputed here:

  variants are ranked by priorD, only the top _ELIGIBLE_VARIANTS may be picked
  pick weight      w_i = 1/(1 + 1.4 i)         i is the rank, 0 is the best
  pick chance      p_i = w_i / sum(w)
  swap chance      s_i = _SwapProb(lam, priorD_i)      lam from AutoLegibility
  drawn as own     d_i = p_i (1 - s_i)
  drawn as anchor  sum(p_i s_i)                the writer's ideal letterform

The join term in the real pick (entry height against where the pen already is)
is left out, so these are the unjoined chances.
"""
import json
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
CP2 = HERE.parent
sys.path.insert(0, str(CP2))   #every pipeline file now sits in CleanPipeline2

import numpy as np                                            # noqa: E402
from docx import Document                                      # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH                  # noqa: E402
from docx.shared import Inches, Pt, RGBColor                   # noqa: E402
from PIL import Image                                          # noqa: E402

import segment
import paths                                                   # noqa: E402
import synthesise as syn                                       # noqa: E402

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

OUT = HERE / "results"
GLY = OUT / "images" / "glyphs"
GLY.mkdir(parents=True, exist_ok=True)
DOCX = OUT / "Glyphs.docx"

PX = 64                  # pixels per x-height unit when drawing a glyph
PEN = 3                  # pen width in pixels
MAX_COLS = 9             # variants shown per table before it wraps
THUMB = 0.34             # inches wide per glyph cell


def Draw(strokes, frame, path):
    """draw one glyph into a shared frame so variants stay comparable in size"""
    x0, y0, x1, y1 = frame
    W = max(10, int((x1 - x0) * PX) + 6)
    H = max(10, int((y1 - y0) * PX) + 6)
    img = np.full((H, W), 255, np.uint8)
    for s in strokes:
        if len(s) < 2:
            continue
        pts = [(float((x - x0) * PX) + 3, float((y1 - y) * PX) + 3)   # y is up
               for x, y in s]
        segment.DrawPolyline(img, pts, 0, PEN)
    base = int((y1 - 0.0) * PX) + 3                                   # baseline row
    if 0 <= base < H:
        img[base, ::3] = np.minimum(img[base, ::3], 205)              # faint baseline
    Image.fromarray(img).save(path)
    return W, H


def Frame(groups):
    """union box over a character's variants, padded, baseline always inside"""
    pts = [p for g in groups for s in g for p in s]
    if not pts:
        return (0.0, 0.0, 1.0, 1.0)
    xs = [p[0] for p in pts]
    ys = [p[1] for p in pts]
    return (min(xs) - 0.08, min(min(ys), -0.05) - 0.08,
            max(xs) + 0.08, max(max(ys), 1.0) + 0.08)


def Probabilities(cands, lam):
    """the selection rule from _GlyphSource, as numbers"""
    ranked = sorted(range(len(cands)), key=lambda i: cands[i].get("priorD", 0.5))
    elig = ranked[:syn._ELIGIBLE_VARIANTS]
    w = np.array([1.0 / (1.0 + 1.4 * i) for i in range(len(elig))])
    p = w / w.sum() if w.sum() > 0 else w
    rows = []
    anchor = 0.0
    for rank, idx in enumerate(ranked):
        d = float(cands[idx].get("priorD", 0.5))
        s = syn._SwapProb(lam, d)
        pick = float(p[rank]) if rank < len(elig) else 0.0
        anchor += pick * s
        rows.append(dict(index=idx, rank=rank + 1, priorD=round(d, 4),
                         eligible=rank < len(elig),
                         pickPct=round(100 * pick, 2),
                         swapPct=round(100 * s, 2),
                         drawnPct=round(100 * pick * (1 - s), 2)))
    return rows, round(100 * anchor, 2)


def Cell(t, i, j, text, bold=False, size=7.5, grey=False):
    c = t.cell(i, j)
    c.text = ""
    par = c.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_after = Pt(0)
    r = par.add_run(text)
    r.bold, r.font.size = bold, Pt(size)
    if grey:
        r.font.color.rgb = RGBColor(0x90, 0x90, 0x90)
    return c


def CharTable(doc, ch, cands, rows, anchorPct, anchorImg, frame, writer):
    """one table per character: the glyphs on top, their chances underneath"""
    order = [r for r in rows]
    for start in range(0, len(order), MAX_COLS):
        chunk = order[start:start + MAX_COLS]
        extra = 1 if (start == 0 and anchorImg is not None) else 0
        t = doc.add_table(rows=6, cols=1 + len(chunk) + extra)
        t.style = "Table Grid"
        labels = ["glyph", "rank by priorD", "prior distance",
                  "pick chance", "swap to anchor", "drawn as this variant"]
        for i, lab in enumerate(labels):
            Cell(t, i, 0, lab, bold=True)
        for j, r in enumerate(chunk, 1):
            g = cands[r["index"]]
            p = GLY / ("%s_%s_%d.png" % (writer, _Safe(ch), r["index"]))
            Draw(g["strokes"], frame, p)
            cell = t.cell(0, j)
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.add_run().add_picture(str(p), width=Inches(THUMB))
            off = not r["eligible"]
            Cell(t, 1, j, str(r["rank"]), grey=off)
            Cell(t, 2, j, "%.3f" % r["priorD"], grey=off)
            Cell(t, 3, j, "never" if off else "%.1f%%" % r["pickPct"], grey=off)
            Cell(t, 4, j, "-" if off else "%.0f%%" % r["swapPct"], grey=off)
            Cell(t, 5, j, "0%" if off else "%.1f%%" % r["drawnPct"], grey=off)
        if extra:
            j = 1 + len(chunk)
            cell = t.cell(0, j)
            cell.text = ""
            par = cell.paragraphs[0]
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.add_run().add_picture(str(anchorImg), width=Inches(THUMB))
            Cell(t, 1, j, "anchor", bold=True)
            Cell(t, 2, j, "-")
            Cell(t, 3, j, "-")
            Cell(t, 4, j, "-")
            Cell(t, 5, j, "%.1f%%" % anchorPct, bold=True)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _Safe(ch):
    return "%02x" % ord(ch)


def Main():
    doc = Document()
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Inches(0.7)
        sec.top_margin = sec.bottom_margin = Inches(0.7)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(9.5)

    doc.add_heading("Stored glyph libraries, and the chance each one is used", 0)
    p = doc.add_paragraph()
    p.add_run("D.R. Braune, 23539195.  Every variant the profile builder kept "
              "for every writer, drawn from the stored stroke coordinates, "
              "beside the probability the synthesiser picks it.  Generated by "
              "testing/glyphs.py").italic = True
    doc.add_paragraph(
        "Variants are ranked by prior distance, which is how far the shape sits "
        "from the pooled shape of that character, so rank 1 is the most typical "
        "letter the writer actually produced. Only the top %d are eligible: "
        "sampling evenly over a dozen variants averages a hand's signature away. "
        "Pick weight falls as 1/(1+1.4 i) with rank i, normalised over the "
        "eligible ones. A picked variant can still be swapped for the writer's "
        "anchor letterform, with probability rising from the legibility lambda "
        "and the variant's own prior distance; at or above 0.55 the swap is "
        "certain. The last row is what actually reaches the page. Rank and "
        "prior distance are stored, every chance is recomputed from the "
        "selection rule, and the join term is left out so these are the "
        "unjoined chances." % syn._ELIGIBLE_VARIANTS)

    store = {}
    for w in paths.AllWriters:
        fp = paths.Profiles / ("%s.json" % w)
        if not fp.exists():
            continue
        prof = json.load(open(fp, encoding="utf-8"))
        lib = prof.get("glyphs", {})
        lam = float(syn.AutoLegibility(prof))
        doc.add_page_break()
        doc.add_heading("Writer %s" % w, 1)
        nVar = sum(len(v) for v in lib.values())
        ds = [float(g.get("priorD", 0.5)) for v in lib.values() for g in v]
        doc.add_paragraph(
            "%d characters, %d stored variants, median prior distance %.3f, "
            "legibility lambda %.2f, connectedness %.3f. %s"
            % (len(lib), nVar, float(np.median(ds)) if ds else 0.0, lam,
               float(prof.get("connectedness", 0.0)),
               ("Lambda is high, so most picks are swapped for the anchor."
                if lam > 0.5 else
                "Lambda is low, so the writer's own shapes are kept almost "
                "always.")))
        store[w] = dict(lam=round(lam, 3), characters=len(lib), variants=nVar,
                        chars={})
        print("\n  %s  %d chars  %d variants  lam %.2f" % (w, len(lib), nVar, lam))
        for ch in sorted(lib):
            cands = lib[ch]
            if not cands:
                continue
            rows, anchorPct = Probabilities(cands, lam)
            core = syn._CoreGlyph(ch, prof)
            frame = Frame([g["strokes"] for g in cands]
                          + ([core[0]] if core else []))
            anchorImg = None
            if core:
                anchorImg = GLY / ("%s_%s_anchor.png" % (w, _Safe(ch)))
                Draw(core[0], frame, anchorImg)
            h = doc.add_paragraph()
            r = h.add_run("'%s'   %d variants" % (ch, len(cands)))
            r.bold, r.font.size = True, Pt(10)
            CharTable(doc, ch, cands, rows, anchorPct, anchorImg, frame, w)
            store[w]["chars"][ch] = dict(variants=len(cands),
                                         anchorPct=anchorPct, rows=rows)
        print("    done, %d characters drawn" % len(store[w]["chars"]))

    doc.save(DOCX)
    (OUT / "glyphs.json").write_text(json.dumps(store, indent=1), encoding="utf-8")
    print("\nwrote %s" % DOCX)
    print("wrote %s" % (OUT / "glyphs.json"))


if __name__ == "__main__":
    Main()
